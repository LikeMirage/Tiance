from __future__ import annotations

import os

import pytest
from app.services.document_conversion.markdown_docx import table_layout
from app.services.document_conversion.markdown_docx.converter import convert_markdown_to_docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from app.services.document_conversion.markdown_docx.table_cell_layout import measure_cell
from app.services.document_conversion.markdown_docx.table_layout import calculate_column_widths
from app.services.document_conversion.markdown_docx.text_measurement import FontTextMeasurer
from app.services.document_conversion.markdown_docx.word_table_layout import apply_column_widths


class DeterministicMeasurer:
    backend = "deterministic-test"

    def measure(self, text: str, *, role: str = "text", bold: bool = False) -> float:
        width = sum(
            10.0 if "\u3400" <= char <= "\u9fff" else 2.5 if char.isspace() else 5.0
            for char in text
        )
        return width * (1.1 if bold else 1.0)


def test_cell_layout_tracks_minimum_and_single_line_preferred_widths() -> None:
    measurer = DeterministicMeasurer()
    sentence = measure_cell("This is a long English sentence for wrapping.", measurer)
    unbreakable = measure_cell("Supercalifragilisticexpialidocious", measurer)

    assert sentence.minimum_width < sentence.preferred_width
    assert unbreakable.minimum_width == unbreakable.preferred_width
    assert sentence.wrap(50).lines > sentence.wrap(150).lines
    assert unbreakable.wrap(50).overflow_points > 0


def test_measured_wrap_cost_gives_long_content_more_space() -> None:
    widths = calculate_column_widths(
        ["ID", "说明", "状态"],
        [
            ["1", "The central limit theorem describes convergence for random variables.", "OK"],
            ["2", "短说明", "OK"],
        ],
        available_width_points=504.0,
        cell_padding_points=10.0,
        measurer=DeterministicMeasurer(),
    )

    assert sum(widths) == pytest.approx(100.0)
    assert widths[1] > widths[0] * 3
    assert widths[1] > widths[2] * 3


def test_equal_columns_remain_equal_after_filling_page_width() -> None:
    count = 30
    widths = calculate_column_widths(
        ["A"] * count,
        [["x"] * count],
        available_width_points=504.0,
        cell_padding_points=10.0,
        measurer=DeterministicMeasurer(),
    )

    assert sum(widths) == pytest.approx(100.0)
    assert max(widths) - min(widths) < 1e-9


def test_word_table_is_explicitly_centered_in_available_page_width() -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)

    apply_column_widths(table, [50.0, 50.0], document)

    alignment = table._tbl.tblPr.find(qn("w:jc"))
    assert table.alignment == WD_TABLE_ALIGNMENT.CENTER
    assert alignment is not None
    assert alignment.get(qn("w:val")) == "center"


def test_formula_exposes_operator_breakpoints_without_losing_minimum_width() -> None:
    cell = measure_cell(
        r"损失函数 $\frac{1}{n}\sum_{i=1}^{n}(y_i-\hat{y}_i)^2$",
        DeterministicMeasurer(),
    )

    assert cell.minimum_width >= 18.0
    assert cell.preferred_width > cell.minimum_width
    assert cell.wrap(cell.minimum_width).lines > 1
    assert cell.wrap(cell.minimum_width / 2).overflow_points > 0


def test_windows_uses_real_gdi_font_measurement() -> None:
    with FontTextMeasurer(
        chinese_font="微软雅黑",
        english_font="Times New Roman",
        math_font="Cambria Math",
        size_points=10.5,
    ) as measurer:
        width = measurer.measure("中文 Hello 123")
        emoji_width = measurer.measure("🚀")
        backend = measurer.backend

    assert width > 0
    assert emoji_width > 0
    if os.name == "nt":
        assert backend == "windows-gdi"


class ScaledMeasurer(DeterministicMeasurer):
    def __init__(self, scale):
        self.scale = scale

    def measure(self, text, *, role="text", bold=False):
        # Different font sizes and font roles must produce separate layouts.
        role_scale = {"text": 1.0, "code": 1.2, "math": 0.9}[role]
        return super().measure(text, role=role, bold=bold) * self.scale * role_scale


MIXED_HEADERS = ["编号", "说明", "公式/代码", "状态"]
MIXED_ROWS = [
    ["1", "**中文粗体**，English sentence with spaces。", r"$x+y=z$ 与 $a\leq b$", "完成"],
    ["2", "换行<br>第二行\n第三行", chr(96) + "unbreakable_identifier_abcdefghij" + chr(96), ""],
    ["3", "[可见链接](https://example.com/very/long/path)", "![图片](local.png)", "待处理"],
    ["4", "（中文标点）《边界》，。", r"$\frac{1}{n}\sum_{i=1}^{n}x_i$", "处理中"],
    ["5", "short"],
    ["6", "", "", "", "多余单元格"],
]


def _uncached_wrap(cell, width, _cache):
    return cell.wrap(width)


@pytest.mark.parametrize("width,padding", [
    (1.0, 10.0), (93.625, 10.375), (223.375, 0.0), (504.0, 10.0), (841.889, 9.875),
])
@pytest.mark.parametrize("scale", [0.8, 1.0, 1.3])
def test_wrap_reuse_keeps_exact_column_widths_across_layouts(monkeypatch, width, padding, scale):
    options = dict(available_width_points=width, cell_padding_points=padding,
                   measurer=ScaledMeasurer(scale))
    actual = calculate_column_widths(MIXED_HEADERS, MIXED_ROWS, **options)
    monkeypatch.setattr(table_layout, "_cached_wrap", _uncached_wrap)
    expected = calculate_column_widths(MIXED_HEADERS, MIXED_ROWS, **options)
    assert actual == expected
    assert sum(actual) == pytest.approx(100.0)


def test_wrap_cache_keeps_fractional_boundaries_and_distinct_font_layouts():
    from math import nextafter

    cache = {}
    plain = measure_cell("AB CD", DeterministicMeasurer())
    bold = measure_cell("AB CD", DeterministicMeasurer(), bold=True)
    boundary = plain.preferred_width
    below = nextafter(boundary, 0.0)
    assert table_layout._cached_wrap(plain, below, cache).lines == 2
    assert table_layout._cached_wrap(plain, boundary, cache).lines == 1
    assert table_layout._cached_wrap(bold, boundary, cache).lines == 2
    assert table_layout._cached_wrap(plain, boundary, cache) == plain.wrap(boundary)


def test_large_table_measures_each_cell_width_once_without_changing_allocation(monkeypatch):
    from collections import Counter

    headers = ["编号", "项目名称", "情况说明", "处理方案", "验收结果", "备注"]
    rows = [[str(i), f"测试项目{i}", f"这是第{i}项的情况说明，需要核对详细内容。",
             f"按步骤处理第{i}项问题并检查处理结果。", f"第{i}项的验收记录和相关说明。",
             f"附加备注{i}，保留完整内容。"] for i in range(200)]
    options = dict(available_width_points=504.0, cell_padding_points=10.0,
                   measurer=DeterministicMeasurer())
    original_wrap = table_layout.CellLayout.wrap
    calls = Counter()

    def counted(cell, width):
        calls[id(cell), width] += 1
        return original_wrap(cell, width)

    monkeypatch.setattr(table_layout.CellLayout, "wrap", counted)
    actual = calculate_column_widths(headers, rows, **options)
    reused_count = sum(calls.values())
    assert max(calls.values()) == 1
    calls.clear()
    monkeypatch.setattr(table_layout, "_cached_wrap", _uncached_wrap)
    expected = calculate_column_widths(headers, rows, **options)
    assert actual == expected
    assert reused_count < sum(calls.values()) * 0.15


def test_wrap_cache_does_not_survive_table_completion_or_failure(monkeypatch):
    # Reuse the exact same layouts across calls: an accidental global cache
    # would skip real measurements in the second conversion.
    cells = table_layout._measure_table(MIXED_HEADERS, MIXED_ROWS, DeterministicMeasurer())
    monkeypatch.setattr(table_layout, "_measure_table", lambda *_: cells)
    original_wrap = table_layout.CellLayout.wrap
    calls = 0
    fail = False

    def counted(cell, width):
        nonlocal calls
        calls += 1
        if fail and calls == 5:
            raise RuntimeError("measurement failed")
        return original_wrap(cell, width)

    monkeypatch.setattr(table_layout.CellLayout, "wrap", counted)
    options = dict(available_width_points=504.0, cell_padding_points=10.0,
                   measurer=DeterministicMeasurer())
    expected = calculate_column_widths(MIXED_HEADERS, MIXED_ROWS, **options)
    count_per_table = calls
    assert count_per_table > 5
    calls = 0
    fail = True
    with pytest.raises(RuntimeError, match="measurement failed"):
        calculate_column_widths(MIXED_HEADERS, MIXED_ROWS, **options)
    fail = False
    calls = 0
    assert calculate_column_widths(MIXED_HEADERS, MIXED_ROWS, **options) == expected
    assert calls == count_per_table


@pytest.mark.parametrize("orientation", ["portrait", "landscape"])
def test_docx_table_xml_and_warnings_match_without_wrap_cache(tmp_path, monkeypatch, orientation):
    from zipfile import ZipFile

    markdown = r"""# 表格回归

| 编号 | 说明 | 数学 |
| --- | --- | --- |
| 1 | **中文粗体** English words with spaces | $x+y=z$ |
| 2 | 第一行<br>第二行，与中文标点。 | CODE_FORMULA |
| 3 | [可见链接](https://example.com/reference) | $a\leq b$ |

第二张表：

| 状态 | 说明 |
| --- | --- |
| 完成 | 内容不同的另一张表，检查列宽缓存隔离。 |
| 待处理 | Long English sentence that should wrap in a narrow column. |
""".replace("CODE_FORMULA", chr(96) + "$x+y=z$" + chr(96))
    cached_path = tmp_path / "cached.docx"
    expected_path = tmp_path / "uncached.docx"
    actual_warnings = convert_markdown_to_docx(
        markdown, cached_path, base_path=tmp_path, page_orientation=orientation,
    )
    monkeypatch.setattr(table_layout, "_cached_wrap", _uncached_wrap)
    expected_warnings = convert_markdown_to_docx(
        markdown, expected_path, base_path=tmp_path, page_orientation=orientation,
    )
    assert actual_warnings == expected_warnings
    actual_doc = Document(cached_path)
    expected_doc = Document(expected_path)
    assert len(actual_doc.tables) == len(expected_doc.tables) == 2
    assert [table._tbl.xml for table in actual_doc.tables] == [
        table._tbl.xml for table in expected_doc.tables
    ]
    with ZipFile(cached_path) as actual, ZipFile(expected_path) as expected:
        assert actual.read("word/document.xml") == expected.read("word/document.xml")

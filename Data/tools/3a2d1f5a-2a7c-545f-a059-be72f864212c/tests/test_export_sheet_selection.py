from __future__ import annotations

import importlib.util
import json
import os
from pathlib import Path
import runpy
import sys
from tempfile import TemporaryDirectory
import types
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


TOOL_ROOT = Path(__file__).resolve().parents[1]
PROGRAM = TOOL_ROOT / "program"
SCHEMA_PATH = TOOL_ROOT / ".tool/input.schema.json"
validator = runpy.run_path(str(TOOL_ROOT.parents[2] / "1_PythonServer/app/services/tools/tool_execution_arguments.py"))["validate_tool_arguments"]
runtime = types.ModuleType("tiance_runtime")
runtime.run_tool = lambda handler: None
spec = importlib.util.spec_from_file_location("excel_selection_main", PROGRAM / "main.py")
MAIN = importlib.util.module_from_spec(spec)
with patch.dict(sys.modules, {"tiance_runtime": runtime}), patch.object(sys, "path", [str(PROGRAM), *sys.path]):
    spec.loader.exec_module(MAIN)


class ExportSheetSelectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        workbook = Workbook()
        workbook.active.title = "Sheet1"
        workbook.active.append(["ONLY_FIRST_SHEET"])
        workbook.create_sheet("Sheet2").append(["ONLY_SECOND_SHEET"])
        workbook.save(self.root / "input.xlsx")
        workbook.close()
        self.schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8-sig"))
        self.args = {"action": "export_markdown", "input_path": "input.xlsx",
                     "content_path": "out/content.md", "format_path": "out/format.md"}

    def call(self, args):
        with patch.dict(os.environ, {"TIANCE_WORKSPACE_ROOT": str(self.root)}):
            return MAIN.run(args)

    def test_schema_has_no_duplicate_keys(self):
        def unique(pairs):
            result = {}
            for key, value in pairs:
                self.assertNotIn(key, result)
                result[key] = value
            return result
        json.loads(SCHEMA_PATH.read_text(encoding="utf-8-sig"), object_pairs_hook=unique)

    def test_selected_sheet_passes_backend_validation_and_excludes_other_sheet(self):
        args = {**self.args, "export_sheets": ["Sheet1"]}
        self.assertEqual(validator(args, self.schema), [])
        result = self.call(args)
        self.assertTrue(result["ok"], result)
        self.assertEqual([item["name"] for item in result["data"]["sheets"]], ["Sheet1"])
        content = (self.root / "out/content.md").read_text(encoding="utf-8")
        formats = (self.root / "out/format.md").read_text(encoding="utf-8")
        self.assertIn("ONLY_FIRST_SHEET", content)
        self.assertNotIn("ONLY_SECOND_SHEET", content)
        self.assertNotIn("Sheet2", formats)

    def test_omitting_selection_still_exports_all_sheets(self):
        result = self.call(self.args)
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["data"]["sheet_count"], 2)

    def test_invalid_selection_never_silently_exports_everything(self):
        for selection in [[], None, "Sheet1", [""], ["  "], [123], [{"name": "Sheet1"}]]:
            with self.subTest(selection=selection):
                result = self.call({**self.args, "export_sheets": selection})
                self.assertFalse(result["ok"])
                self.assertEqual(result["error_info"]["code"], "INVALID_ARGUMENT")
                self.assertFalse((self.root / "out/content.md").exists())

    def test_legacy_sheets_field_returns_clear_migration_error(self):
        result = self.call({**self.args, "sheets": ["Sheet1"]})
        self.assertFalse(result["ok"])
        self.assertIn("export_sheets", result["error"])
        self.assertFalse((self.root / "out/content.md").exists())

    def test_unknown_sheet_fails_without_exporting_all(self):
        result = self.call({**self.args, "export_sheets": ["Missing"]})
        self.assertFalse(result["ok"])
        self.assertFalse((self.root / "out/content.md").exists())

    def test_create_keeps_sheets_objects(self):
        args = {"action": "create", "output_path": "created.xlsx",
                "sheets": [{"name": "Made", "data_blocks": [{"start_cell": "A1", "rows": [["value"]]}]}]}
        self.assertEqual(validator(args, self.schema), [])
        result = self.call(args)
        self.assertTrue(result["ok"], result)
        workbook = load_workbook(self.root / "created.xlsx")
        try:
            self.assertEqual(workbook.sheetnames, ["Made"])
            self.assertEqual(workbook["Made"]["A1"].value, "value")
        finally:
            workbook.close()

    def test_published_examples_match_schema(self):
        examples = json.loads((TOOL_ROOT / ".tool/examples.json").read_text(encoding="utf-8-sig"))
        for example in examples:
            with self.subTest(title=example["title"]):
                self.assertEqual(validator(json.loads(example["content"]), self.schema), [])


if __name__ == "__main__":
    unittest.main()
